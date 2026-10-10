"""Workbook provenance helpers shared by import and audit tooling."""
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import load_workbook

REVISED4_DATASET = "revised4"
REVISED4_SCHEMA_VERSION = "revised4-6-sheet-2026-10-09"
REVISED4_SHEETS = (
    "Participant_Data",
    "Prediction",
    "BIA_MNA_Staff",
    "Satisfaction",
    "Data_Dictionary ",
    "Model_Logic",
)
SHEET_ROLES = {
    "Participant_Data": "observations",
    "Prediction": "derived_results",
    "BIA_MNA_Staff": "observations",
    "Satisfaction": "observations",
    "Data_Dictionary ": "definition",
    "Model_Logic": "algorithm",
}


def workbook_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def workbook_manifest(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False, read_only=True)
    if tuple(workbook.sheetnames) != REVISED4_SHEETS:
        raise RuntimeError(f"unexpected revised4 sheets: {workbook.sheetnames}")
    manifest = {}
    for sheet_name in REVISED4_SHEETS:
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        formula_cells = 0
        nonempty_cells = 0
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    nonempty_cells += 1
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells += 1
        manifest[sheet_name] = {
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "headers": headers,
            "formula_cells": formula_cells,
            "nonempty_cells": nonempty_cells,
        }
    return manifest


def _json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def workbook_sheet_snapshots(path: Path) -> list[dict]:
    values_book = load_workbook(path, data_only=True, read_only=True)
    formulas_book = load_workbook(path, data_only=False, read_only=True)
    snapshots = []
    for sheet_name in REVISED4_SHEETS:
        values_sheet = values_book[sheet_name]
        formulas_sheet = formulas_book[sheet_name]
        values = [
            [_json_value(cell.value) for cell in row]
            for row in values_sheet.iter_rows()
        ]
        formulas = {}
        for row in formulas_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas[cell.coordinate] = cell.value
        canonical = json.dumps(
            {"values": values, "formulas": formulas},
            ensure_ascii=False, sort_keys=True, separators=(",", ":"),
        )
        snapshots.append({
            "sheet_name": sheet_name,
            "sheet_role": SHEET_ROLES[sheet_name],
            "row_count": values_sheet.max_row,
            "column_count": values_sheet.max_column,
            "content_sha256": sha256(canonical.encode("utf-8")).hexdigest(),
            "values_json": values,
            "formulas_json": formulas,
        })
    return snapshots
