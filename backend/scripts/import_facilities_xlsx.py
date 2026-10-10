"""Upsert Bangkok public health centers from the approved four-column workbook."""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.database import SessionLocal  # noqa: E402
from app.models import Facility  # noqa: E402


EXPECTED_HEADERS = ("ศูนย์บริการสาธารณสุข", "ที่ตั้ง", "Google Map", "เบอร์โทร")


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u200b", "")).strip()


def import_facilities(path: Path, apply: bool = False, replace: bool = False) -> tuple[int, int, int]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = tuple(clean(value) for value in next(rows))
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"Unexpected headers: {headers!r}")

    created = updated = deactivated = 0
    imported_names: set[str] = set()
    with SessionLocal() as db:
        for order, row in enumerate(rows, start=1):
            name, address, map_link, phone = (clean(value) for value in row[:4])
            if not any((name, address, map_link, phone)):
                continue
            if not all((name, address, map_link, phone)):
                raise ValueError(f"Incomplete facility row {order + 1}")
            imported_names.add(name)
            record = db.query(Facility).filter(Facility.name == name).first()
            if record is None:
                number = re.search(r"ศูนย์บริการสาธารณสุข\s*(\d+)", name)
                record = Facility(name=name, code=f"BKK-HC-{int(number.group(1)):02d}" if number else None)
                db.add(record)
                created += 1
            else:
                updated += 1
            record.address = address
            record.map_link = map_link
            record.phone = phone
            record.type = "ศูนย์บริการสาธารณสุข"
            record.display_order = order
            record.is_active = True
            record.is_deleted = False
        if replace:
            extras = db.query(Facility).filter(Facility.name.notin_(imported_names), Facility.is_deleted.is_(False)).all()
            for record in extras:
                record.is_active = False
                record.is_deleted = True
            deactivated = len(extras)
        if apply:
            db.commit()
        else:
            db.rollback()
    return created, updated, deactivated


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--apply", action="store_true", help="commit changes; default is dry-run")
    parser.add_argument("--replace", action="store_true", help="soft-delete facilities absent from workbook")
    args = parser.parse_args()
    created, updated, deactivated = import_facilities(args.workbook, args.apply, args.replace)
    print(f"{'APPLIED' if args.apply else 'DRY-RUN'}: created={created}, updated={updated}, deactivated={deactivated}")


if __name__ == "__main__":
    main()
