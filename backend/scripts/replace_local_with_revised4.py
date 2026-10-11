"""Replace the local participant cohort with the 200 observed revised4 rows.

Dry-run is default. This intentionally does not invent raw option labels, visit
dates, assessment IDs or staff IDs that are absent from the workbook.
"""
import argparse
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app.database import SessionLocal
from app.models import (BIARecord, MNAResponse, Respondent, SatisfactionResponse,
                        SANSAResponse, ScoringRuleVersion, SourceImportBatch,
                        SourceImportSheet, Visit)
from app.services.sansa_definition import DB_SCORING_VERSION
from app.services.sansa_prediction_storage import save_sansa_prediction
from app.services.assessment_revision_service import append_assessment_revision
from app.services.source_provenance_service import (
    REVISED4_DATASET,
    REVISED4_SCHEMA_VERSION,
    workbook_manifest,
    workbook_sheet_snapshots,
    workbook_sha256,
)

BOOK = Path(__file__).resolve().parents[2] / "Use SANSA Integrated WebApp Data (revised4).xlsx"
SOURCE = "revised4"


def rows(ws):
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    return [dict(zip(header, (cell.value for cell in row))) for row in ws.iter_rows(min_row=2, max_row=201)]


def dec(value):
    return None if value is None else Decimal(str(value))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    wb = load_workbook(BOOK, data_only=True, read_only=True)
    participant = rows(wb["Participant_Data"])
    prediction = rows(wb["Prediction"])
    staff = rows(wb["BIA_MNA_Staff"])
    satisfaction = rows(wb["Satisfaction"])
    ids = [row["record_id"] for row in participant]
    if len(ids) != 200 or len(set(ids)) != 200 or set(ids) != set(range(1, 201)):
        raise RuntimeError("revised4 populated cohort must be exactly record_id 1..200")
    for dataset in (prediction, staff, satisfaction):
        if [row["record_id"] for row in dataset] != ids:
            raise RuntimeError("record_id order differs between revised4 sheets")
    for p, pred in zip(participant, prediction):
        if p["D1_total"] != p["ass_meal"] + p["ass_intake"] + p["ass_type"]:
            raise RuntimeError(f"D1 mismatch record {p['record_id']}")
        if p["D2_total"] != sum(p[k] for k in ("ass_rice_starch","ass_meat_egg","ass_milk","ass_fruit","ass_veget","ass_water")):
            raise RuntimeError(f"D2 mismatch record {p['record_id']}")
        if p["D3_total"] != sum(p[k] for k in ("ass_sweetdrink","ass_cook","ass_oil")):
            raise RuntimeError(f"D3 mismatch record {p['record_id']}")
        if (p["D2_total"], p["D3_total"], p["wc_cm"]) != (pred["D2_total"], pred["D3_total"], pred["wc_cm"]):
            raise RuntimeError(f"Prediction inputs mismatch record {p['record_id']}")
    print({"mode": "apply" if args.apply else "dry-run", "validated_rows": 200,
           "assessment_dates_present": sum(row["assessment_date"] is not None for row in participant)})
    if not args.apply:
        return

    db = SessionLocal()
    try:
        sansa_version = db.query(ScoringRuleVersion).filter_by(
            instrument_name="SANSA", version_number=DB_SCORING_VERSION, is_active=True).one()
        mna_version = db.query(ScoringRuleVersion).filter_by(instrument_name="MNA", is_active=True).first()
        if mna_version is None:
            raise RuntimeError("active MNA scoring version is required")
        for respondent in db.query(Respondent).all():
            db.delete(respondent)
        db.flush()
        digest = workbook_sha256(BOOK)
        batch = db.query(SourceImportBatch).filter_by(source_sha256=digest).first()
        if batch is None:
            batch = SourceImportBatch(
                source_dataset=REVISED4_DATASET,
                source_filename=BOOK.name,
                source_sha256=digest,
                schema_version=REVISED4_SCHEMA_VERSION,
                populated_rows=200,
                sheet_manifest=workbook_manifest(BOOK),
            )
            db.add(batch)
            db.flush()
        if db.query(SourceImportSheet).filter_by(import_batch_id=batch.id).count() == 0:
            for snapshot in workbook_sheet_snapshots(BOOK):
                db.add(SourceImportSheet(import_batch_id=batch.id, **snapshot))
        sex_map = {1: "male", 2: "female"}; edu = {1:"no_education",2:"primary",3:"secondary",4:"vocational",5:"bachelor",6:"graduate"}
        marital = {1:"single",2:"married",3:"widowed",4:"divorced"}; income = {1:"under_10k",2:"10k_20k",3:"20k_30k",4:"over_30k"}
        living = {1:"alone",2:"with_spouse",3:"with_friends",4:"with_family"}
        score_cols = ["scr_bw","scr_app","scr_act","scr_dx","ass_meal","ass_intake","ass_type","ass_rice_starch","ass_meat_egg","ass_milk","ass_fruit","ass_veget","ass_water","ass_sweetdrink","ass_cook","ass_oil"]
        mna_scores = [f"mna_s{i}" for i in range(1,7)] + [f"mna_a{i}" for i in range(1,13)]
        for p, pred, b, sat in zip(participant, prediction, staff, satisfaction):
            diseases = {"diseases": [k.lower() for k in ("DM","HT","DLP","CKD","CA") if p[k] == 1], "other": p["Other_text"] if p["Other"] == 1 else "", "other_flag": int(p["Other"]), "source_other_text": p["Other_text"] or ""}
            sources = [name for col,name in (("Income_oldage","old_age"),("Income_Pension","pension"),("Income_spouse","spouse"),("Income_children","children"),("Income_work","work")) if p[col] == 1]
            person = Respondent(respondent_code=f"R4-{int(p['record_id']):04d}", source_dataset=REVISED4_DATASET, source_record_id=int(p["record_id"]), source_import_batch_id=batch.id, status="elderly", age=p["age"], sex=sex_map[p["sex"]], education_level=edu[p["Education"]], marital_status=marital[p["Marital_status"]], monthly_income=income[p["Income"]], income_sources=sources, chronic_diseases=diseases, living_arrangement=living[p["Living_with"]])
            db.add(person); db.flush()
            visit = Visit(respondent_id=person.id, visit_number=1, visit_date=p["assessment_date"], source_assessment_id=p["assessment_id"], visit_type="baseline")
            db.add(visit); db.flush()
            scores = {f"q{i}_score": dec(p[col]) for i,col in enumerate(score_cols, 1)}
            sansa = SANSAResponse(visit_id=visit.id, scoring_version_id=sansa_version.id, **scores,
                screening_total=dec(sum(p[c] for c in score_cols[:4])), diet_total=dec(sum(p[c] for c in score_cols[4:])), total_score=dec(sum(p[c] for c in score_cols)), result_level=None,
                source_dataset=SOURCE, answer_capture_mode="score_only")
            db.add(sansa)
            bia = BIARecord(visit_id=visit.id, age=p["age"], sex=sex_map[p["sex"]], waist_circumference_cm=dec(p["wc_cm"]), height_cm=dec(b["height_cm"]), weight_kg=dec(b["weight_kg"]), bmi=dec(b["bmi"]), fat_mass_kg=dec(b["fat_mass_kg"]), body_fat_percentage=dec(b["pbf"]), visceral_fat_level=dec(b["vfat_level"]), muscle_mass_kg=dec(b["muscle_mass_kg"]), bone_mass_kg=dec(b["bone_mass_kg"]), water_percentage=dec(b["tbw"]), metabolic_rate=b["bmr_kcal"], mac_cm=dec(b["MAC"]), cc_cm=dec(b["CC"]), source_staff_id=b["staff_id"], measurement_date=b["assessment_date"])
            db.add(bia); db.flush()
            mna = MNAResponse(visit_id=visit.id, scoring_version_id=mna_version.id,
                source_dataset=SOURCE, answer_capture_mode="score_only", entry_mode="staff",
                mna_screen_total=dec(b["mna_screen_total"]), mna_ass_total=dec(b["mna_ass_total"]),
                mna_total=dec(b["mna_total"]), result_category=b["mna_total_cat"],
                mac_cm=dec(b["MAC"]), cc_cm=dec(b["CC"]),
                **{key:dec(b[key]) for key in mna_scores})
            db.add(mna)
            db.flush()
            append_assessment_revision(db, "BIA", bia, action_type="import")
            append_assessment_revision(db, "MNA", mna, action_type="import")
            db.add(SatisfactionResponse(visit_id=visit.id, q1_clarity=sat["sat_clarity"], q2_ease_of_use=sat["sat_easy"], q3_confidence=sat["sat_confident"], q4_presentation=sat["sat_design"], q5_results_display=sat["sat_result"], q6_usefulness=sat["sat_benefit"], q7_overall_satisfaction=sat["sat_overall"], comments=sat["sat_comment"]))
            save_sansa_prediction(db, visit.id, d2_total=p["D2_total"], d3_total=p["D3_total"], waist_cm=p["wc_cm"], age_group01=p["age_group01"])
        db.commit()
        print({"created": 200, "respondents": db.query(Respondent).count(), "visits": db.query(Visit).count(), "sansa": db.query(SANSAResponse).count(), "bia": db.query(BIARecord).count(), "mna": db.query(MNAResponse).count(), "satisfaction": db.query(SatisfactionResponse).count()})
    except Exception:
        db.rollback(); raise
    finally:
        db.close()


if __name__ == "__main__":
    main()
